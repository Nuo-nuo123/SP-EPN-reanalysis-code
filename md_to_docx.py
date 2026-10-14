# -*- coding: utf-8 -*-
"""Convert the manuscript markdown to a Word document (with real tables)."""
import re, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import sys
MD = sys.argv[1] if len(sys.argv) > 1 else r'D:\claude-study\Spinal Ependymoma\Clinical_Epigenetics_初稿.md'
OUT = sys.argv[2] if len(sys.argv) > 2 else r'D:\claude-study\Spinal Ependymoma\Clinical_Epigenetics_初稿.docx'

text = open(MD, encoding='utf-8').read()

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

def _title_border_black():
    """Keep the Title style's bottom rule but make it black; all other heading
    styles stay without a rule."""
    st = doc.styles['Title']
    pPr = st.element.get_or_add_pPr()
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is None:
        return
    bottom = pBdr.find(qn('w:bottom'))
    if bottom is None:
        return
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), '000000')
    if bottom.get(qn('w:themeColor')) is not None:
        del bottom.attrib[qn('w:themeColor')]

_title_border_black()

def add_para(text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.bold = True
        else:
            r = p.add_run(part)
            if italic:
                r.italic = True
    return p

def set_three_line_borders(tbl):
    """Three-line (booktabs) style: 1.5pt top rule, 0.75pt rule under the header row,
    1.5pt bottom rule; no vertical or interior borders."""
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'bottom'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '12')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    for edge in ('left', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'none')
        borders.append(el)
    tblPr.append(borders)
    for cell in tbl.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tb = OxmlElement('w:tcBorders')
        b = OxmlElement('w:bottom')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '6')
        b.set(qn('w:color'), '000000')
        tb.append(b)
        tcPr.append(tb)

def add_markdown_table(rows):
    """rows: list of lists (first = header). Renders a three-line Word table.
    Trailing empty cells in a row are merged into the preceding cell."""
    for r in rows:
        while len(r) > 1 and not r[-1].strip():
            r.pop()
    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = tbl.cell(i, j)
            txt = row[j].strip() if j < len(row) else ''
            bold = txt.startswith('**') and txt.endswith('**') and len(txt) > 4
            if bold:
                txt = txt[2:-2]
            cell.text = ''
            p = cell.paragraphs[0]
            r = p.add_run(txt)
            r.font.size = Pt(9)
            r.font.name = 'Times New Roman'
            if i == 0 or bold:
                r.bold = True
    for i, row in enumerate(rows):
        if len(row) < ncols:
            tbl.cell(i, len(row) - 1).merge(tbl.cell(i, ncols - 1))
    # vertical merge of consecutive identical first-column cells (e.g., repeated
    # "Gene region" / "CpG island relation" labels in Table 2)
    for i in range(1, len(rows)):
        a = rows[i - 1][0].strip() if len(rows[i - 1]) else ''
        b = rows[i][0].strip() if len(rows[i]) else ''
        if a and a == b and len(rows[i - 1]) == ncols and len(rows[i]) == ncols:
            tbl.cell(i - 1, 0).merge(tbl.cell(i, 0))
    set_three_line_borders(tbl)
    return tbl

def parse_pipe_table(lines, start_idx):
    """Parse markdown pipe table starting at start_idx; return (rows, next_idx).
    Splits only on unescaped pipes; '\\|' inside cells renders as a literal '|'."""
    rows = []
    i = start_idx
    while i < len(lines) and lines[i].strip().startswith('|'):
        line = lines[i].strip()
        if re.match(r'^[\s:\-\|]+$', line.replace(' ', '')):  # separator row
            i += 1
            continue
        cells = [c.strip().replace('\\|', '|')
                 for c in re.split(r'(?<!\\)\|', line.strip('|'))]
        rows.append(cells)
        i += 1
    return rows, i

lines = text.split('\n')
i = 0
while i < len(lines):
    line = lines[i].rstrip()

    if line.strip() == '---':
        i += 1
        continue

    if line.startswith('# '):
        p = doc.add_heading(line[2:].strip(), level=0)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)
        i += 1
        continue

    m = re.match(r'^## (.+)$', line)
    if m:
        title = m.group(1).strip()
        if title == 'Abstract':
            doc.add_heading('Abstract', level=1)
        elif title == 'References':
            doc.add_page_break()
            doc.add_heading('References', level=1)
        elif title == 'Abbreviations':
            doc.add_heading('Abbreviations', level=1)
        elif title == 'Declarations':
            doc.add_heading('Declarations', level=1)
        else:
            doc.add_heading(title, level=1)
        i += 1
        continue

    if line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=2)
        i += 1
        continue

    m = re.match(r'^\*\*(Manuscript type|Target journal|Status):\*\* (.*)$', line)
    if m:
        add_para(m.group(2), bold_prefix=m.group(1) + ': ', italic=True)
        i += 1
        continue

    # Table/figure caption line
    m = re.match(r'^\*\*(Table \d+|Fig\. S?\d+[A-D]?)\*\* (.*)$', line)
    if m:
        add_para(m.group(2), bold_prefix=m.group(1) + ' ')
        i += 1
        continue

    # figure image reference
    m = re.match(r'^!\[.*?\]\((.+?)\)$', line)
    if m:
        img = m.group(1)
        if not os.path.isabs(img):
            img = os.path.join(os.path.dirname(os.path.abspath(MD)), img.replace('/', os.sep))
        if os.path.exists(img):
            try:
                from PIL import Image
                w_px, _ = Image.open(img).size
                width_in = min(6.2, w_px / 300.0)
            except Exception:
                width_in = 6.2
            doc.add_picture(img, width=Inches(width_in))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            print('missing image:', img)
        i += 1
        continue

    m = re.match(r'^(Background|Methods|Results|Conclusions|Keywords|Availability of data and materials|Competing interests|Ethics approval and consent to participate|Consent for publication|Funding|Acknowledgements|Authors\' contributions):\s?(.*)$', line)
    if m:
        add_para(m.group(2), bold_prefix=m.group(1) + ': ')
        i += 1
        continue

    m = re.match(r'^\*\*([^*]+)\*\*$', line)
    if m:
        p = doc.add_paragraph()
        r = p.add_run(m.group(1))
        r.bold = True
        i += 1
        continue

    if re.match(r'^\d+\. ', line):
        p = doc.add_paragraph(line)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        i += 1
        continue

    # markdown pipe table
    if line.strip().startswith('|'):
        rows, ni = parse_pipe_table(lines, i)
        add_markdown_table(rows)
        i = ni
        continue

    if not line.strip():
        i += 1
        continue

    add_para(line.strip())
    i += 1

try:
    doc.save(OUT)
    print('saved:', OUT)
except PermissionError:
    tmp = OUT.replace('.docx', '_tmp.docx')
    doc.save(tmp)
    print('LOCKED (Word open); saved to:', tmp)
